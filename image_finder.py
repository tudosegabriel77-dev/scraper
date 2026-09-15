import os
import re
import time
import html
from io import BytesIO
from urllib.parse import unquote, urlparse

import pandas as pd
import requests
import urllib3

from PIL import Image, UnidentifiedImageError

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D

from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

import logging

log = logging.getLogger("image_finder")

# ============================================================
# SETTINGS
# ============================================================

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}

VERIFY_SSL = False

if not VERIFY_SSL:
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

SIZE_PATTERN = re.compile(
    r"\b(?:XS|S|M|L|XL|XXL|XXXL|2XL|3XL)\b",
    re.IGNORECASE
)

SIZE_RANGE_PATTERN = re.compile(
    r"\b(?:XS|S|M|L|XL|XXL|XXXL|2XL|3XL)"
    r"\s*-\s*"
    r"(?:XS|S|M|L|XL|XXL|XXXL|2XL|3XL)\b",
    re.IGNORECASE
)

# ============================================================
# SESSION
# ============================================================

def make_session():
    session = requests.Session()
    session.headers.update(HEADERS)

    retries = Retry(
        total=2,
        backoff_factor=0.4,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET", "HEAD"],
    )
    adapter = HTTPAdapter(
        max_retries=retries,
        pool_connections=20,
        pool_maxsize=20
    )
    session.mount("http://", adapter)
    session.mount("https://", adapter)

    # Warm up Bing so consent cookies are set for the proxy IP
    try:
        session.get("https://www.bing.com/", timeout=20, verify=VERIFY_SSL)
        log.info("Bing homepage warm-up done")
    except Exception as e:
        log.warning("Bing warm-up failed: %s", e)

    return session

# ============================================================
# TEXT / QUERY HELPERS
# ============================================================

def normalize_spaces(text):
    return re.sub(r"\s+", " ", str(text)).strip()

def remove_sizes_only(text):
    t = str(text)
    t = SIZE_RANGE_PATTERN.sub("", t)
    t = SIZE_PATTERN.sub("", t)
    return normalize_spaces(t)

def clean_product_identity(text):
    t = remove_sizes_only(text)
    t = re.sub(r"[-,;/]+", " ", t)
    t = normalize_spaces(t)
    return t

def build_cache_key(query):
    return clean_product_identity(query).lower()

def build_query_variants(query):
    exact_query = normalize_spaces(query)
    no_size_query = remove_sizes_only(query)
    cleaned_query = clean_product_identity(query)
    variants = []
    for q in [exact_query, no_size_query, cleaned_query]:
        q = normalize_spaces(q)
        if q and q not in variants:
            variants.append(q)
    return variants

# ============================================================
# URL HELPERS
# ============================================================

def normalize_url(url):
    if not url:
        return None
    url = html.unescape(str(url).strip())
    url = url.replace("\\/", "/")
    url = unquote(url)
    url = url.strip("'\" ,;<>")
    if not url.startswith(("http://", "https://")):
        return None
    return url

def build_request_headers_for_url(url):
    headers = dict(HEADERS)
    try:
        parsed = urlparse(url)
        origin = f"{parsed.scheme}://{parsed.netloc}"
        headers["Referer"] = origin + "/"
    except Exception:
        pass
    return headers

# ============================================================
# BING IMAGE SEARCH
# ============================================================

def search_bing_images(query, session, max_results=12):
    try:
        r = session.get(
            "https://www.bing.com/images/search",
            params={"q": query},
            timeout=20,
            allow_redirects=True,
            headers=HEADERS,
            verify=VERIFY_SSL,
        )

        log.info("  Bing HTTP %s | bytes=%d", r.status_code, len(r.text))

        if r.status_code != 200:
            return []

        matches = re.findall(r'murl&quot;:&quot;(.*?)&quot;', r.text)
        if not matches:
            matches = re.findall(r'"murl":"(.*?)"', r.text)

        log.info("  Regex found %d raw murl matches", len(matches))

        urls = []
        for m in matches:
            u = normalize_url(m)
            if u and u not in urls:
                urls.append(u)
            if len(urls) >= max_results:
                break

        return urls
    except Exception as e:
        log.warning("  search_bing_images exception: %s", e)
        return []

def find_candidate_urls_bing(query, session, log_func=None):
    def _log(msg):
        if log_func:
            log_func(msg)
        log.info(msg)

    all_candidates = []
    variants = build_query_variants(query)
    for q in variants:
        _log(f" Trying query: {q}")
        candidates = search_bing_images(q, session=session, max_results=12)
        if not candidates:
            _log(" No results")
            continue
        for idx, url in enumerate(candidates, start=1):
            short_url = url[:120] + ("..." if len(url) > 120 else "")
            _log(f" Candidate {idx}: {short_url}")
            if url not in all_candidates:
                all_candidates.append(url)
    return all_candidates

# ============================================================
# IMAGE DOWNLOAD
# ============================================================

def download_image_to_png(url, out_path, session, log_func=None):
    def _log(msg):
        if log_func:
            log_func(msg)
        log.info(msg)

    try:
        headers = build_request_headers_for_url(url)
        r = session.get(
            url,
            timeout=25,
            allow_redirects=True,
            headers=headers,
            verify=VERIFY_SSL
        )
        _log(f" HTTP {r.status_code} | content-type={r.headers.get('Content-Type')}")
        if r.status_code != 200:
            return False
        img = Image.open(BytesIO(r.content))
        img.load()
        img = img.convert("RGB")
        img.save(out_path, format="PNG")
        return True
    except UnidentifiedImageError as e:
        _log(f" PIL could not identify image: {e}")
        return False
    except requests.exceptions.SSLError as e:
        _log(f" SSL error: {e}")
        return False
    except requests.exceptions.RequestException as e:
        _log(f" Request error: {e}")
        return False
    except Exception as e:
        _log(f" Other error: {type(e).__name__}: {e}")
        return False

# ============================================================
# IMAGE RESIZE
# ============================================================

def resize_image_keep_ratio(img_path, max_w=144, max_h=135):
    with Image.open(img_path) as img:
        w, h = img.size
        ratio = min(max_w / w, max_h / h)
        if ratio < 1:
            new_w = max(1, int(w * ratio))
            new_h = max(1, int(h * ratio))
            img = img.resize((new_w, new_h))
            img.save(img_path)

# ============================================================
# WRITE EXCEL
# ============================================================

def write_excel_with_embedded_images(excel_path, rows_data, log_func=None):
    def _log(msg):
        if log_func:
            log_func(msg)
        log.info(msg)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    CELL_WIDTH_PX = 160
    CELL_HEIGHT_PX = 150
    IMG_WIDTH_PX = int(CELL_WIDTH_PX * 0.9)
    IMG_HEIGHT_PX = int(CELL_HEIGHT_PX * 0.9)

    ws["A1"] = "Reference"
    ws["B1"] = "Link"
    ws["C1"] = "Image"

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 23

    for i, row in enumerate(rows_data, start=2):
        ref = row.get("Reference", "")
        chosen_url = row.get("ChosenURL", "")
        local_img_path = row.get("LocalImagePath", "")

        ws[f"A{i}"] = ref
        ws[f"B{i}"] = chosen_url
        ws[f"A{i}"].alignment = Alignment(vertical="top")
        ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"C{i}"].alignment = Alignment(vertical="top")
        ws.row_dimensions[i].height = CELL_HEIGHT_PX * 0.75

        if not local_img_path or not os.path.exists(local_img_path):
            continue

        try:
            resize_image_keep_ratio(local_img_path, max_w=IMG_WIDTH_PX, max_h=IMG_HEIGHT_PX)
            with Image.open(local_img_path) as pil_img:
                actual_w, actual_h = pil_img.size
            xl_img = XLImage(local_img_path)
            xl_img.width = actual_w
            xl_img.height = actual_h
            offset_x = max(0, int((CELL_WIDTH_PX - actual_w) / 2))
            offset_y = max(0, int((CELL_HEIGHT_PX - actual_h) / 2))
            marker = AnchorMarker(
                col=2,
                colOff=offset_x * 9525,
                row=i - 1,
                rowOff=offset_y * 9525
            )
            size = XDRPositiveSize2D(cx=actual_w * 9525, cy=actual_h * 9525)
            xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
            ws.add_image(xl_img)
            _log(f"Embedded image in row {i}")
        except Exception as e:
            _log(f"Could not embed image in row {i}: {e}")

    wb.save(excel_path)

# ============================================================
# MAIN EXCEL PROCESSING
# ============================================================

def process_excel(input_file, output_dir, log_func=None):
    def _log(msg):
        if log_func:
            log_func(msg)
        log.info(msg)

    if not os.path.exists(input_file):
        raise FileNotFoundError(f"Input file not found: {input_file}")

    os.makedirs(output_dir, exist_ok=True)

    base_name = os.path.splitext(os.path.basename(input_file))[0]
    output_excel = os.path.join(output_dir, f"{base_name}_with_images.xlsx")
    images_dir = os.path.join(output_dir, f"{base_name}_images")
    os.makedirs(images_dir, exist_ok=True)

    _log("Reading Excel file...")
    df_raw = pd.read_excel(input_file, header=None)
    if df_raw.empty:
        raise ValueError("The Excel file is empty.")

    references = df_raw.iloc[:, 0].fillna("").astype(str).str.strip()
    references = references[references != ""]
    if references.empty:
        raise ValueError("Column A does not contain any valid values.")

    refs_list = references.tolist()
    total = len(refs_list)
    _log(f"Processing {total} article(s)...")

    session = make_session()
    success_cache = {}
    rows_data = []

    for i, ref in enumerate(refs_list, start=1):
        _log(f"[{i}/{total}] Searching: {ref}")
        cache_key = build_cache_key(ref)

        if cache_key in success_cache:
            cached = success_cache[cache_key]
            _log(f" Reused successful cached image for: {cache_key}")
            rows_data.append({
                "Reference": ref,
                "ChosenURL": cached["url"],
                "LocalImagePath": cached["local_path"],
            })
            continue

        candidates = find_candidate_urls_bing(ref, session=session, log_func=_log)
        chosen_url = ""
        local_img_path = ""

        if candidates:
            _log(" Trying candidate downloads...")
            for idx, url in enumerate(candidates, start=1):
                short_url = url[:120] + ("..." if len(url) > 120 else "")
                _log(f" Download candidate {idx}: {short_url}")
                temp_img_path = os.path.join(images_dir, f"row_{i}.png")
                if download_image_to_png(url, temp_img_path, session, log_func=_log):
                    chosen_url = url
                    local_img_path = temp_img_path
                    _log(" Accepted and downloaded")
                    success_cache[cache_key] = {
                        "url": chosen_url,
                        "local_path": local_img_path,
                    }
                    break
                else:
                    _log(" Failed")
        else:
            _log(" No candidates found")

        rows_data.append({
            "Reference": ref,
            "ChosenURL": chosen_url,
            "LocalImagePath": local_img_path,
        })
        time.sleep(0.05)

    _log("Writing Excel file and embedding images...")
    write_excel_with_embedded_images(output_excel, rows_data, log_func=_log)
    _log(f"Done: {output_excel}")
    return output_excel
